# -*- coding: utf-8 -*-
"""Fills an arbitrary .xlsx template whose cells contain ``{{ }}`` placeholders.

This is plain Python with no Odoo/ORM dependency on purpose, so it can be
unit-tested (and reasoned about) on its own, independently of which models
the values it receives came from. Turning a bracket's tag/category *names*
into ids is Odoo's job (see ``models/metrics/base.py``) -- this file only
ever deals in plain strings and cell references.

Conventions (see doc/template_authoring.md for the full writer-facing guide)
-----------------------------------------------------------------------------
- ``{{some_key}}`` alone in a cell is replaced by the raw value registered
  under ``some_key`` in the ``metrics`` dict passed to :meth:`fill` -- so a
  number/date stays a number/date and any ``SUM()`` elsewhere in the sheet
  keeps working.
- ``{{some_key}}`` mixed with other text (``"Total : {{total}} EUR"``) does a
  text substitution instead, and the cell stays text.
- ``{{rows:block_name}}`` alone as a cell's own value marks that *row* as a
  repeating block: ``metrics[block_name]`` must then be a list of dicts, one
  per line to produce. Every other placeholder on that same row is resolved
  against each dict in turn, once per repetition. A dotted placeholder's
  first segment is just a readability alias and is otherwise ignored --
  ``{{move.date}}`` and ``{{date}}`` both resolve to ``record["date"]``.
- ``{{key["Name1","Name2"]}}`` scopes that one placeholder's own metric call
  to just those tag names -- the *only* way a metric ever gets a tag/PoS
  filter, since the wizard/CLI only ever supply a date range.
  ``{{key[tag=[...], pos_categ=[...]]}}`` scopes on both a tag and a
  PoS-category axis at once. ``{{key["Name"].field}}`` reads one field
  out of a list-returning ("detail") metric's single matching record,
  without a row. ``{{key["Name1","Name2"].field}}`` (more than one name)
  turns its *row* into a repeating block sized and ordered by that name
  list, exactly like ``{{rows:...}}`` but driven by names typed here
  instead of "however many the metric returns". ``{{key[].field}}`` (no
  names at all) *is* "however many the metric returns" -- equivalent to an
  explicit ``{{rows:key}}`` marker, just spelled with ``.field`` on every
  cell of the row instead of a separate marker cell and bare
  ``{{field}}``s. A bare, unquoted, spreadsheet-style token (``A1``) stands
  for that same-sheet cell's own text, split on commas -- so the names only
  ever need to be typed once.

Known limitations:
- Expanding a rows-block (explicit ``{{rows:...}}`` or an implicit
  ``.field`` one, named or not) inserts worksheet rows via openpyxl's
  ``insert_rows``, which does not reliably renumber merged cells or
  absolute-row formulas that sit *below* the block on the same sheet. Put
  rows-blocks at the bottom of a sheet, or on their own sheet, to stay safe.
- A placeholder driving an implicit (``.field`` with zero or 2+ names) row
  expansion must be a cell's entire content -- it can't be mixed with other
  text like a single-name ``.field`` or a plain ``{{key}}`` can.
- A cell reference only ever points at a cell on the *same* sheet.
"""

import copy
import csv
import io
import re
import zipfile
from collections import namedtuple

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

_PLACEHOLDER = re.compile(r"\{\{\s*(.+?)\s*\}\}")
_ROWS_BLOCK = re.compile(r"^rows:\s*(.+)$", re.S)
_FIELD_NAME = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
_CELL_REF = re.compile(r"[A-Za-z]+[0-9]+")
_AXIS_ASSIGN = re.compile(r"^([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*)$", re.S)

#: The only filter axes a bracket can name -- kept in sync with
#: ``models/metrics/base.py``'s ``_AXIS_MODELS``/``_AXIS_PARAM``, which map
#: each of these to an Odoo model/field and a metric kwarg respectively.
VALID_AXES = ("tag", "pos_categ")

ParsedPlaceholder = namedtuple("ParsedPlaceholder", "base_key axes field")
#: One distinct metric call a template needs -- ``axes`` is ``{}`` for no
#: filter, else ``{axis_name: [names]}``. ``sheet``/``row`` are only for
#: error messages (the first place this exact call was seen).
MetricCall = namedtuple("MetricCall", "canonical_key base_key axes sheet row")


class TemplateError(Exception):
    """A template file or its placeholders aren't usable as given."""


def _open_workbook(template_bytes):
    """``load_workbook``, translating "this isn't a real .xlsx" failures
    into a message someone without a Python traceback can act on.

    The most common real-world cause: a file saved from LibreOffice Calc
    with the format dropdown still on "ODF Spreadsheet (.ods)" even though
    its name ends in .xlsx. It's a valid zip archive either way, so the
    error only shows up once openpyxl looks for the Excel-specific parts
    inside it -- surfacing as a zip/KeyError, not an obviously-file-format
    complaint.
    """
    try:
        return load_workbook(io.BytesIO(template_bytes))
    except (KeyError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise TemplateError(
            "This file doesn't look like a real .xlsx workbook. If you "
            "built it in LibreOffice Calc, check the format dropdown in "
            "the Save dialog is set to 'Excel 2007-365 (.xlsx)', not "
            "'ODF Spreadsheet (.ods)' -- a file only named .xlsx isn't "
            "enough. (%s: %s)" % (type(exc).__name__, exc)
        ) from exc


def canonical_key(base_key, axes):
    """The dict key a given (base_key, axes) call is stored/looked up under.

    Unchanged (``base_key`` itself) when there's no filter, so every
    existing template's ``metrics`` dict shape keeps working untouched.
    Shared, exact contract with ``models/metrics/base.py``'s
    ``compute_metrics`` -- both build this the same way so a lookup always
    matches, and two placeholders asking for the same filter compute once.
    """
    if not axes:
        return base_key
    parts = [
        "%s=%s" % (axis, ",".join(axes[axis]))
        for axis in sorted(axes)
    ]
    return "%s[%s]" % (base_key, ";".join(parts))


def max_axis_len(axes):
    return max((len(names) for names in axes.values()), default=0)


def driving_axis(axes):
    """The one axis (if any) with more than one name -- the axis whose
    names dictate an implicit ``.field`` row expansion's row count/order.

    Raises if more than one axis has multiple names: which one should
    drive the row count would be undefined.
    """
    driving = [(axis, names) for axis, names in axes.items() if len(names) > 1]
    if len(driving) > 1:
        raise TemplateError(
            "more than one filter axis has multiple names (%s) -- only one "
            "axis can drive a '.field' row expansion"
            % ", ".join(axis for axis, _ in driving)
        )
    return driving[0] if driving else None


# -- placeholder grammar (pure string/spreadsheet parsing, no Odoo) --------

def _split_key_bracket_field(text):
    """``"key[bracket].field"`` -> ``(key, bracket_or_None, field_or_None)``."""
    start = text.find("[")
    if start == -1:
        return text.strip(), None, None
    base_key = text[:start].strip()
    depth = 0
    end = None
    for i in range(start, len(text)):
        if text[i] == "[":
            depth += 1
        elif text[i] == "]":
            depth -= 1
            if depth == 0:
                end = i
                break
    if end is None:
        raise TemplateError("unbalanced '[' in placeholder '%s'" % text)
    bracket_content = text[start + 1:end]
    tail = text[end + 1:].strip()
    field = None
    if tail:
        if not tail.startswith("."):
            raise TemplateError("unexpected text after ']' in placeholder '%s'" % text)
        field = tail[1:].strip()
        if not _FIELD_NAME.fullmatch(field):
            raise TemplateError("invalid field name '%s' in placeholder '%s'" % (field, text))
    return base_key, bracket_content, field


def _split_top_level(text):
    """Split on commas that aren't inside ``"..."`` or ``[...]``."""
    parts = []
    depth = 0
    in_quotes = False
    current = []
    for ch in text:
        if ch == '"':
            in_quotes = not in_quotes
            current.append(ch)
        elif ch == "[" and not in_quotes:
            depth += 1
            current.append(ch)
        elif ch == "]" and not in_quotes:
            depth -= 1
            current.append(ch)
        elif ch == "," and depth == 0 and not in_quotes:
            parts.append("".join(current))
            current = []
        else:
            current.append(ch)
    parts.append("".join(current))
    return [p.strip() for p in parts if p.strip()]


def _parse_literal_list(inner):
    inner = inner.strip()
    if not inner:
        return []
    try:
        names = next(csv.reader([inner], skipinitialspace=True))
    except csv.Error as exc:
        raise TemplateError("malformed name list '%s': %s" % (inner, exc)) from exc
    names = [n.strip() for n in names]
    if any(not n for n in names):
        raise TemplateError("empty name in list '%s'" % inner)
    return names


def _resolve_cell_ref(ref, sheet):
    cell = sheet[ref.upper()]
    value = cell.value
    if isinstance(value, str) and "{{" in value:
        raise TemplateError(
            "cell reference '%s' (sheet '%s') points at another placeholder"
            % (ref, sheet.title)
        )
    if value is None or (isinstance(value, str) and not value.strip()):
        raise TemplateError("cell reference '%s' (sheet '%s') is empty" % (ref, sheet.title))
    names = [n.strip() for n in str(value).split(",") if n.strip()]
    if not names:
        raise TemplateError("cell reference '%s' (sheet '%s') has no names" % (ref, sheet.title))
    return names


def _parse_value(text, sheet):
    """One axis's (or the bracket's, for the no-axis sugar form) value:
    a bracketed literal list, a bare cell reference, or a single literal."""
    text = text.strip()
    if text.startswith("[") and text.endswith("]"):
        return _parse_literal_list(text[1:-1])
    if _CELL_REF.fullmatch(text):
        return _resolve_cell_ref(text, sheet)
    return _parse_literal_list(text)


def _parse_axes(bracket_content, sheet):
    content = bracket_content.strip()
    if not content:
        return {}
    top_parts = _split_top_level(content)
    if _AXIS_ASSIGN.match(top_parts[0]):
        axes = {}
        for part in top_parts:
            match = _AXIS_ASSIGN.match(part)
            if not match:
                raise TemplateError("malformed filter axis '%s'" % part)
            axis, value_text = match.group(1), match.group(2)
            if axis not in VALID_AXES:
                raise TemplateError(
                    "unknown filter axis '%s' -- valid axes are: %s"
                    % (axis, ", ".join(VALID_AXES))
                )
            if axis in axes:
                raise TemplateError("filter axis '%s' given more than once" % axis)
            names = _parse_value(value_text, sheet)
            if not names:
                raise TemplateError("filter axis '%s' has no names" % axis)
            axes[axis] = names
        return axes
    names = _parse_value(content, sheet)
    return {"tag": names} if names else {}


def _parse_placeholder(token, sheet):
    base_key, bracket_content, field = _split_key_bracket_field(token)
    if not base_key:
        raise TemplateError("placeholder is missing a key: '%s'" % token)
    axes = {} if bracket_content is None else _parse_axes(bracket_content, sheet)
    return ParsedPlaceholder(base_key, axes, field)


class TemplateFiller:
    """Fills one .xlsx template's placeholders with computed values."""

    def __init__(self, template_bytes):
        self._template_bytes = template_bytes

    def fill(self, metrics):
        """Return the filled workbook as bytes.

        :param metrics: dict mapping a :func:`canonical_key` to either a
            scalar (for a plain ``{{key}}`` placeholder) or a list of dicts
            (for a matching ``{{rows:key}}`` block, or a ``.field``
            accessor).
        """
        workbook = _open_workbook(self._template_bytes)
        for sheet in workbook.worksheets:
            self._fill_sheet(sheet, metrics)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    # -- sheet level ----------------------------------------------------

    def _fill_sheet(self, sheet, metrics):
        blocks = self._find_row_blocks(sheet)
        # Expand bottom-up so row numbers found above stay valid as rows
        # get inserted below them.
        for block in sorted(blocks, key=lambda b: b[1], reverse=True):
            kind, row_number = block[0], block[1]
            if kind == "explicit":
                _, _, anchor_col, parsed = block
                canon = canonical_key(parsed.base_key, parsed.axes)
                if canon not in metrics:
                    raise TemplateError(
                        "Template references unknown block '%s' "
                        "(sheet '%s', row %s)" % (canon, sheet.title, row_number)
                    )
                self._expand_explicit_row_block(sheet, row_number, anchor_col, metrics[canon])
            else:
                _, _, canon = block
                self._expand_implicit_row_block(sheet, row_number, canon, metrics)

        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = self._substitute_metrics(cell.value, sheet, metrics)

    def _find_row_blocks(self, sheet):
        explicit = {}  # row -> (anchor_col, ParsedPlaceholder)
        implicit = {}  # row -> canonical_key
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                whole = _PLACEHOLDER.fullmatch(cell.value.strip())
                if not whole:
                    continue
                token = whole.group(1)
                rows_match = _ROWS_BLOCK.match(token)
                if rows_match:
                    parsed = _parse_placeholder(rows_match.group(1), sheet)
                    if parsed.field is not None:
                        raise TemplateError(
                            "'{{rows:...}}' cannot use '.field' (sheet '%s', "
                            "row %s) -- give the repeating row's other cells "
                            "their own bare {{field}} placeholders instead"
                            % (sheet.title, cell.row)
                        )
                    if cell.row in explicit or cell.row in implicit:
                        raise TemplateError(
                            "row %s has more than one repeating-row spec "
                            "(sheet '%s')" % (cell.row, sheet.title)
                        )
                    explicit[cell.row] = (cell.column, parsed)
                    continue

                parsed = _parse_placeholder(token, sheet)
                if parsed.field is None or max_axis_len(parsed.axes) == 1:
                    continue
                # No names at all (like {{rows:...}}, but spelled with
                # ".field" -- "however many the metric returns") or more
                # than one name (a fixed, named set) both drive a repeating
                # row; :func:`driving_axis` only needs to validate the
                # latter ("at most one axis can name several entries").
                driving_axis(parsed.axes)
                canon = canonical_key(parsed.base_key, parsed.axes)
                if cell.row in explicit:
                    raise TemplateError(
                        "row %s mixes an explicit '{{rows:...}}' marker with "
                        "'%s[...].%s' (sheet '%s')"
                        % (cell.row, parsed.base_key, parsed.field, sheet.title)
                    )
                existing = implicit.get(cell.row)
                if existing is None:
                    implicit[cell.row] = canon
                elif existing != canon:
                    raise TemplateError(
                        "row %s has inconsistent filters between its "
                        "'.field' placeholders (sheet '%s')" % (cell.row, sheet.title)
                    )

        blocks = [
            ("explicit", row_number, anchor_col, parsed)
            for row_number, (anchor_col, parsed) in explicit.items()
        ]
        blocks += [
            ("implicit", row_number, canon)
            for row_number, canon in implicit.items()
        ]
        return blocks

    def _expand_explicit_row_block(self, sheet, anchor_row, anchor_col, records):
        max_col = sheet.max_column
        # The anchor cell only ever carries the {{rows:...}} marker -- clear
        # it before capturing the template row. It can sit in any column
        # (not necessarily the row's first one), so a field placeholder can
        # occupy column A instead.
        sheet.cell(row=anchor_row, column=anchor_col).value = None

        if not records:
            for col in range(1, max_col + 1):
                cell = sheet.cell(row=anchor_row, column=col)
                if isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = None
            return

        template_values = [
            sheet.cell(row=anchor_row, column=col).value
            for col in range(1, max_col + 1)
        ]

        extra_rows = len(records) - 1
        if extra_rows > 0:
            sheet.insert_rows(anchor_row + 1, amount=extra_rows)
            for offset in range(1, extra_rows + 1):
                self._clone_row_style(sheet, anchor_row, anchor_row + offset)

        for offset, record in enumerate(records):
            target_row = anchor_row + offset
            for col, value in enumerate(template_values, start=1):
                if isinstance(value, str) and "{{" in value:
                    sheet.cell(row=target_row, column=col).value = self._substitute_record(
                        value, record
                    )

    def _expand_implicit_row_block(self, sheet, anchor_row, canon, metrics):
        records = self._lookup_metric(canon, metrics)
        if not isinstance(records, list):
            raise TemplateError(
                "'%s' isn't a row-returning metric -- can't drive a "
                "repeating row (sheet '%s', row %s)" % (canon, sheet.title, anchor_row)
            )

        max_col = sheet.max_column
        template_values = [
            sheet.cell(row=anchor_row, column=col).value
            for col in range(1, max_col + 1)
        ]

        count = len(records)
        if count == 0:
            for col in range(1, max_col + 1):
                cell = sheet.cell(row=anchor_row, column=col)
                if isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = None
            return

        extra_rows = count - 1
        if extra_rows > 0:
            sheet.insert_rows(anchor_row + 1, amount=extra_rows)
            for offset in range(1, extra_rows + 1):
                self._clone_row_style(sheet, anchor_row, anchor_row + offset)

        for index in range(count):
            target_row = anchor_row + index
            for col, value in enumerate(template_values, start=1):
                if isinstance(value, str) and "{{" in value:
                    sheet.cell(row=target_row, column=col).value = self._substitute_metrics(
                        value, sheet, metrics, row_index=index
                    )

    @staticmethod
    def _clone_row_style(sheet, source_row, target_row):
        for col in range(1, sheet.max_column + 1):
            source_cell = sheet.cell(row=source_row, column=col)
            target_cell = sheet.cell(row=target_row, column=col)
            if source_cell.has_style:
                target_cell._style = copy.copy(source_cell._style)
        source_dim = sheet.row_dimensions.get(source_row)
        if source_dim is not None and source_dim.height is not None:
            sheet.row_dimensions[target_row].height = source_dim.height

    # -- placeholder resolution ------------------------------------------

    def _substitute_record(self, text, record):
        """Resolve a rows-block's per-record bare {{field}} placeholders."""
        stripped = text.strip()
        whole = _PLACEHOLDER.fullmatch(stripped)
        if whole:
            return self._lookup_field(whole.group(1), record)

        def substitute(match):
            value = self._lookup_field(match.group(1), record)
            return "" if value is None else str(value)

        return _PLACEHOLDER.sub(substitute, text)

    def _substitute_metrics(self, text, sheet, metrics, row_index=None):
        """Resolve any placeholder against the top-level ``metrics`` dict --
        plain keys, bracket-filtered calls, and ``.field`` accessors alike."""
        stripped = text.strip()
        whole = _PLACEHOLDER.fullmatch(stripped)
        if whole:
            return self._resolve_one(whole.group(1), sheet, metrics, row_index)

        def substitute(match):
            value = self._resolve_one(match.group(1), sheet, metrics, row_index)
            return "" if value is None else str(value)

        return _PLACEHOLDER.sub(substitute, text)

    def _resolve_one(self, token, sheet, metrics, row_index):
        parsed = _parse_placeholder(token, sheet)
        canon = canonical_key(parsed.base_key, parsed.axes)
        value = self._lookup_metric(canon, metrics)
        if parsed.field is None:
            return value

        if not isinstance(value, list):
            raise TemplateError(
                "'%s' isn't a row-returning metric -- '.%s' needs one"
                % (parsed.base_key, parsed.field)
            )
        if max_axis_len(parsed.axes) == 1:
            # Exactly one name: a plain single-value substitution, no row.
            if not value:
                return None
            if len(value) > 1:
                raise TemplateError(
                    "'%s[...]' returned more than one row for a single name"
                    % parsed.base_key
                )
            return self._lookup_field(parsed.field, value[0])

        # No name at all (like {{rows:...}}, "however many the metric
        # returns") or more than one name (a fixed, named set): both drive
        # a repeating row, indexed positionally.
        if row_index is None:
            raise TemplateError(
                "'%s[...].%s' isn't driving a repeating row (sheet '%s') "
                "-- put '.%s' on every cell of its own row, nothing else"
                % (parsed.base_key, parsed.field, sheet.title, parsed.field)
            )
        if row_index >= len(value):
            return None
        return self._lookup_field(parsed.field, value[row_index])

    @staticmethod
    def _lookup_metric(key, metrics):
        if key not in metrics:
            raise TemplateError("Template references unknown metric '%s'" % key)
        return metrics[key]

    @staticmethod
    def _lookup_field(placeholder, record):
        # A dotted placeholder's first segment is a readability alias only:
        # {{move.date}} and {{date}} both mean record["date"].
        path = placeholder.split(".", 1)[-1] if "." in placeholder else placeholder
        current = record
        for part in path.split("."):
            if not isinstance(current, dict) or part not in current:
                raise TemplateError("Row has no field '%s'" % placeholder)
            current = current[part]
        return current

    # -- introspection -----------------------------------------------------

    @classmethod
    def required_keys(cls, template_bytes):
        """Return ``(metric_keys, block_keys)`` referenced by a template,
        without computing anything -- lets a caller validate a template or
        show what data it needs before actually running it.

        ``metric_keys`` are plain top-level ``{{key}}``/``{{key[...]}}``
        placeholders' base keys; ``block_keys`` are ``{{rows:key}}`` block
        base keys. Per-record field placeholders living on a rows-block's
        own row are excluded from ``metric_keys`` since they aren't
        top-level metrics.
        """
        workbook = _open_workbook(template_bytes)
        metric_keys = set()
        block_keys = set()
        for sheet in workbook.worksheets:
            block_rows = set()
            tokens_by_row = {}
            for row in sheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    for match in _PLACEHOLDER.finditer(cell.value):
                        token = match.group(1)
                        rows_match = _ROWS_BLOCK.match(token)
                        if rows_match:
                            parsed = _parse_placeholder(rows_match.group(1), sheet)
                            block_keys.add(parsed.base_key)
                            block_rows.add(cell.row)
                        else:
                            parsed = _parse_placeholder(token, sheet)
                            tokens_by_row.setdefault(cell.row, []).append(parsed.base_key)
            for row_number, keys in tokens_by_row.items():
                if row_number not in block_rows:
                    metric_keys.update(keys)
        return metric_keys, block_keys

    @classmethod
    def required_calls(cls, template_bytes):
        """Every distinct metric call (base key + resolved filter axes) this
        template needs, deduplicated by :func:`canonical_key` -- what
        ``compute_metrics``/``validate_calls`` (``models/metrics/base.py``)
        actually iterate over."""
        workbook = _open_workbook(template_bytes)
        calls = {}
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    for match in _PLACEHOLDER.finditer(cell.value):
                        token = match.group(1)
                        rows_match = _ROWS_BLOCK.match(token)
                        inner = rows_match.group(1) if rows_match else token
                        parsed = _parse_placeholder(inner, sheet)
                        canon = canonical_key(parsed.base_key, parsed.axes)
                        if canon not in calls:
                            calls[canon] = MetricCall(
                                canon, parsed.base_key, parsed.axes,
                                sheet.title, cell.row,
                            )
        return list(calls.values())
